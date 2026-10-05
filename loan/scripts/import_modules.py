import json,re,hashlib,warnings,gzip,base64,sys
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import openpyxl
warnings.filterwarnings('ignore',category=UserWarning)
ROOT=Path('loan'); OUT=ROOT/'catalog';OUT.mkdir(exist_ok=True)
stamp=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds')
classification=[
('上階額度','額度管理','集團(1G)、子集團(2S)、客戶(3C)建檔、維護與查詢',['001','005']),
('放_01','額度管理','額度(4F)及產品(5P)建檔、維護、查詢',['001','002','003','004','006','007','008','009','010','011','014']),
('放_02','擔保品管理','擔保品建檔、維護與查詢',['012','013','015','016','017']),
('放_03','擔保品管理','保證人建檔、維護與查詢',['018']),
('放_04','帳戶管理','帳戶建檔、維護與查詢',['019','020','029','031','034','035','036','037','038','039','040','041','042','043']),
('放_05','利息處理','提息與計息',['021','023','030']),
('放_06','利息處理','撥提、補收及補退',['027']),
('放_07','費用處理','代墊、收費、撥提、補收及補退',['027']),
('放_08','帳務處理','還款',['022','025','026']),
('放_09','帳務處理','結清',['028']),
('放_10','帳務處理','調整（含減免）',['028','025']),
('放_11','貸後管理','變更及展延',['021','022','029','033','047']),
('放_12','不良債權','延滯與逾期',['024']),
('放_13','不良債權','催收與呆帳',['044','046','045']),
('放_14','債務協商','債清條例',['032'])]
groups=[dict(id=k,business=b,process=p,ssr=['B2-'+s for s in ss],note='圖片第一列未填放_XX，保留未編號分類；不建立正式情境編號。' if k=='上階額度' else '依使用者提供的分類圖片') for k,b,p,ss in classification]
sources=[]; cases=[]; requirements=[]
def value(v):
 if v is None:return ''
 if isinstance(v,datetime):return v.isoformat()
 if hasattr(v,'text'):return v.text or ''
 return str(v)
def rows(p):
 wf=openpyxl.load_workbook(p,read_only=True,data_only=False);wv=openpyxl.load_workbook(p,read_only=True,data_only=True)
 for s in wf:
  if p.name.startswith('需求清單') and not s.title.startswith('需求清單'):continue
  if p.name.startswith('測試資料表') and not s.title.startswith('測試資料表'):continue
  if p.name.startswith('L7') and '測試案例內容' not in s.title:continue
  yield s.title,list(s.values),list(wv[s.title].values)
def fieldset(headers,raw,cache):
 out=[]
 for i,v in enumerate(raw):
  if v is None:continue
  h=value(headers[i]).strip() if i<len(headers) else ''
  if not h:h='未命名欄位 '+openpyxl.utils.get_column_letter(i+1)
  c=value(cache[i]) if i<len(cache) else ''
  t=value(v)
  f=t if t.startswith('=') else ''
  out.append(dict(column=openpyxl.utils.get_column_letter(i+1),label=h,value=c if f and c else t,formula=f,cache_missing=bool(f and not c)))
 return out
files=sorted(Path(sys.argv[1] if len(sys.argv)>1 else 'module_sources').rglob('*.xlsx'))
for p in files:
 if not (p.name.startswith('需求清單') or p.name.startswith('測試資料表') or p.name.startswith('L7_放')):continue
 kind='requirements' if p.name.startswith('需求清單') else 'l7' if p.name.startswith('L7') else 'l6'
 code=re.search(r'B2-?(\d{3})',p.name) or re.search(r'B2(\d{3})',p.name)
 code=code.group(1) if code else ('043' if 'B2043' in p.name else '032' if 'B2032' in p.name else '')
 group_match=re.search(r'放_(\d+)',p.name);group='放_'+group_match.group(1).zfill(2) if group_match else ''
 sid=hashlib.sha256(p.name.encode()).hexdigest()[:12]; count=0
 for sn,raws,caches in rows(p):
  if kind=='requirements' and not sn.startswith('需求清單'):continue
  if kind=='l6' and not sn.startswith('測試資料表'):continue
  if kind=='l7' and '測試案例內容' not in sn:continue
  header=2 if kind in ('requirements','l7') else 3
  hs=raws[header]
  for i in range(header+1,len(raws)):
   raw=raws[i];cache=caches[i];fs=fieldset(hs,raw,cache)
   def get(label):return next((x['value'] for x in fs if x['label'].replace('\n','').strip()==label),'')
   if kind=='requirements':
    # Source requirement sheets use a real identifier or its Excel formula in A.
    rid=value(cache[0]) or value(raw[0]);desc=value(cache[7]) or value(raw[7]) if len(raw)>7 else ''
    if not rid or rid=='範例' or not desc or (not rid.startswith('=') and not re.search(r'(?:SSR|FSD).*_\d+|B2-\d+.*_\d+',rid)):continue
    if rid.startswith('='):
     # Keep unresolved formula visible; never fabricate a Requirement ID.
     rid='未取得需求編號快取：'+sid+':'+str(i+1)
    requirements.append(dict(id='REQ-'+sid+'-'+str(i+1),key=rid,code='B2-'+code,title=desc,fields=fs,source=p.name,sheet=sn,row=i+1,source_id=sid,kind=kind,groups=[g['id'] for g in groups if 'B2-'+code in g['ssr']],status='原始需求；有效性待逐條確認'))
   else:
    key=get('L6 情境編號') if kind=='l7' else get('Key_情境編號')
    if kind=='l6':
     num=get('情境編號');title=value(cache[2]) or value(raw[2])
     if not num or not title or not re.fullmatch(r'\d+(?:[_-]\d+)*(?:\.0)?',num):continue
     derived=not key or key.startswith('=')
     if derived:key=group+'_'+('_'.join(part.zfill(4) if n==0 else part for n,part in enumerate(re.sub(r'\.0$','',num).replace('-','_').split('_'))))
    else:
     title=get('測案內容說明');derived=False
     if not key or not re.search(r'放[_-]?\d+',key):continue
    cases.append(dict(id=kind.upper()+'-'+sid+'-'+str(i+1),key=key,title=title,group=group,groups=[group],kind=kind,fields=fs,source=p.name,sheet=sn,row=i+1,source_id=sid,key_derived=derived,requirement=get('需求清單編號') or get('SSR'),ticket=get('需求單號') or get('requirements'),expected=get('預期結果補充說明') or get('Expected result'),action=get('Action'),input=get('Input'),status='原檔轉錄；未重新驗收'))
   count+=1
 sources.append(dict(id=sid,name=p.name,kind=kind,count=count,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),imported_at=stamp,group=group,code='B2-'+code if code else ''))
# Latest 032/043 explicitly supplied updated files replace the older sheet per topic.
latest={c:next((s['id'] for s in sources if s['kind']=='requirements' and s['code']=='B2-'+c and '2026' in s['name'] and not s['name'].startswith('需求清單-')),None) for c in ['032','043']}
for r in requirements:r['version_status']='已取代' if r['code'][3:] in latest and latest[r['code'][3:]] and r['source_id']!=latest[r['code'][3:]] else '本次收錄版本'
payloads={}
for k in ['l6','l7','requirements']:
 xs=requirements if k=='requirements' else [c for c in cases if c['kind']==k]
 for n in range(0,len(xs),100):
  name=f'{k}-{n//100:03d}.json';raw=json.dumps(xs[n:n+100],ensure_ascii=False,separators=(',',':')).encode();compressed=base64.b64encode(gzip.compress(raw,compresslevel=9,mtime=0)).decode();(OUT/name).write_text(json.dumps({'encoding':'gzip-base64','payload':compressed},separators=(',',':')));payloads.setdefault(k,[]).append('catalog/'+name)
manifest=dict(updated_at=stamp,classification=groups,sources=sources,datasets=payloads,counts={k:len(requirements) if k=='requirements' else sum(c['kind']==k for c in cases) for k in ['l6','l7','requirements']},scope_note='本次匯入可取得的完整原檔：L6放_04、放_09、放_14；L7放_01。其他分類缺原始案例檔，不以分類表或摘要補造案例。L7按步驟列收錄，同一L6編號可有多筆步驟。圖片第一列未填編號，網站保留為未編號的上階額度分類。')
(ROOT/'modules.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest['counts'],ensure_ascii=False));print('sources',len(sources),'files',sum(map(len,payloads.values())))
